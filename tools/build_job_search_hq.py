"""Build the one-workbook job search HQ (Pipeline, Recruiters CRM, Inbox Log, Job Leads).

Usage:
    python tools/build_job_search_hq.py --recruiters data/recruiters.csv \
        --data data/data.json --out "Damian - Job Search HQ.xlsx"

    # Blank template (no personal data):
    python tools/build_job_search_hq.py --out template.xlsx

Inputs live in data/ (gitignored) because they hold recruiter contact details.
- recruiters.csv: CSV export of "LinkedIn Recruiter Tracker - Damian Ruiz".
- data.json: curated pipeline, job leads, and older message log (see STRATEGY.md).
"""

import argparse
import csv
import json
import re
from collections import OrderedDict
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

STAGES = [
    "1 New - triage",
    "2 Reply owed (me)",
    "3 Waiting on recruiter",
    "4 Submitted / RTR",
    "5 Interviewing",
    "6 Offer",
    "7 Accepted",
    "Nurture",
    "Closed - I declined",
    "Closed - Rejected",
    "Closed - Went cold",
    "Closed - Below floor/off-target",
    "Spam / scam",
]
PRIORITIES = ["Hot", "Warm", "Cold"]
TIERS = ["A", "B", "C", "DNC"]
PAY_TYPES = ["Hourly", "Salary", "C2C/1099", ""]
SOURCES = ["LinkedIn", "Email", "Job board", "Referral", "Phone"]

FLOOR = 85000
TARGET_ROLE = re.compile(
    r"contract|procure|buyer|sourcing|supply|subcontract|vendor|category|purchas|conformance",
    re.I,
)

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="1F3864")
H2_FONT = Font(bold=True, size=12, color="1F3864")
WRAP = Alignment(wrap_text=True, vertical="top")
OVERDUE_FILL = PatternFill("solid", fgColor="F8CBAD")
HOT_FILL = PatternFill("solid", fgColor="FFE699")
CLOSED_FONT = Font(color="808080")


def normalize_status(raw, when, as_of):
    """Map ~50 free-text statuses onto the fixed stage list."""
    s = (raw or "").lower()
    if "spam" in s or "scam" in s:
        return "Spam / scam"
    if "declin" in s:
        return "Closed - I declined"
    if "not a fit" in s or "off-target" in s or "not pursued" in s:
        return "Closed - Below floor/off-target"
    if "offer" in s:
        return "6 Offer"
    if "interview" in s:
        stage = "5 Interviewing"
    elif "rtr" in s or "sent resume" in s or "call booked" in s or "submitted" in s:
        stage = "4 Submitted / RTR"
    elif "connected" in s:
        return "Nurture"
    elif "draft" in s or "not yet contacted" in s or "wants to schedule" in s:
        stage = "2 Reply owed (me)"
    elif any(k in s for k in ("replied", "asked", "interested", "shared phone", "open -", "active")):
        stage = "3 Waiting on recruiter"
    else:
        stage = "1 New - triage"
    # Anything open and untouched for 30+ days is effectively dead.
    try:
        age = (as_of - date.fromisoformat(when)).days
    except ValueError:
        age = 999
    return "Closed - Went cold" if age > 30 else stage


def style_header(ws, headers, widths):
    ws.append(headers)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 30


def add_list_validation(ws, col, list_range, rows=1000):
    dv = DataValidation(type="list", formula1=list_range, allow_blank=True)
    dv.add(f"{col}2:{col}{rows}")
    ws.add_data_validation(dv)


def as_date(v):
    try:
        return date.fromisoformat(v) if v else None
    except ValueError:
        return None


def build(recruiters_csv, data_path, out, as_of, include_log=True, log_url=""):
    data = json.load(open(data_path)) if data_path else {}
    log_rows = list(csv.DictReader(open(recruiters_csv))) if recruiters_csv else []

    wb = Workbook()
    start = wb.active
    start.title = "Start Here"
    pipe = wb.create_sheet("Pipeline")
    recs = wb.create_sheet("Recruiters")
    log = wb.create_sheet("Inbox Log")
    jobs = wb.create_sheet("Job Leads")
    tmpl = wb.create_sheet("Templates")
    lists = wb.create_sheet("Lists")

    # ---- Lists (dropdown sources) ----
    for col, (name, values) in enumerate(
        [("Stage", STAGES), ("Priority", PRIORITIES), ("Tier", TIERS),
         ("Pay type", PAY_TYPES), ("Source", SOURCES)], start=1):
        lists.cell(row=1, column=col, value=name).font = Font(bold=True)
        for r, v in enumerate(values, start=2):
            lists.cell(row=r, column=col, value=v)
    rng = {
        "stage": f"Lists!$A$2:$A${len(STAGES) + 1}",
        "priority": f"Lists!$B$2:$B${len(PRIORITIES) + 1}",
        "tier": f"Lists!$C$2:$C${len(TIERS) + 1}",
        "paytype": f"Lists!$D$2:$D${len(PAY_TYPES) + 1}",
        "source": f"Lists!$E$2:$E${len(SOURCES) + 1}",
    }

    # ---- Pipeline: one row per live opportunity ----
    headers = ["ID", "Stage", "Priority", "Company / end client", "Role", "Recruiter",
               "Agency", "Work model", "Pay (as stated)", "Pay type", "Rate or base",
               "Annualized $", "Meets $85K?", "Source", "First contact", "Last touch",
               "Days idle", "Next action", "Due", "Contact email", "Notes"]
    widths = [6, 22, 9, 28, 34, 22, 20, 14, 22, 10, 10, 13, 11, 10, 12, 12, 9, 60, 12, 30, 45]
    style_header(pipe, headers, widths)
    for i, p in enumerate(data.get("pipeline", []), start=1):
        r = i + 1
        pipe.append([
            f"P{i:03d}", p["stage"], p["priority"], p["company"], p["role"], p["recruiter"],
            p["agency"], p["model"], p["pay"], p["pay_type"], p["annual"] or None,
            f'=IF(K{r}="","",IF(J{r}="Hourly",K{r}*2080,K{r}))',
            f'=IF(L{r}="","?",IF(L{r}>={FLOOR},"Yes","No"))',
            p["source"], as_date(p["first"]), as_date(p["last"]),
            f'=IF(P{r}="","",TODAY()-P{r})',
            p["next"], as_date(p["due"]), p["email"], p["notes"],
        ])
    last = max(pipe.max_row, 2)
    for row in pipe.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.alignment = WRAP
        for idx in (14, 15, 18):
            row[idx].number_format = "yyyy-mm-dd"
        row[11].number_format = "$#,##0"
    open_row = 'LEFT($B2,6)<>"Closed"'
    pipe.conditional_formatting.add(
        "A2:U1000", FormulaRule(formula=[f'AND($S2<>"",$S2<TODAY(),{open_row})'], fill=OVERDUE_FILL))
    pipe.conditional_formatting.add(
        "A2:U1000", FormulaRule(formula=[f'AND($C2="Hot",{open_row})'], fill=HOT_FILL))
    pipe.conditional_formatting.add(
        "A2:U1000", FormulaRule(formula=['LEFT($B2,6)="Closed"'], font=CLOSED_FONT))
    add_list_validation(pipe, "B", rng["stage"])
    add_list_validation(pipe, "C", rng["priority"])
    add_list_validation(pipe, "J", rng["paytype"])
    add_list_validation(pipe, "N", rng["source"])
    pipe.auto_filter.ref = "A1:U1000"

    # ---- Inbox Log: every recruiter message, never deleted ----
    headers = ["Date", "Recruiter", "Agency", "End client", "Role pitched", "Location / model",
               "Email", "Stage (normalized)", "Original status", "Summary", "Notes", "Origin"]
    widths = [11, 22, 24, 24, 32, 26, 30, 22, 30, 60, 60, 16]
    style_header(log, headers, widths)
    for x in log_rows:
        log.append([
            as_date(x["Date Contacted"]), x["Recruiter Name"], x["Company/Agency"],
            x["Recruiting For (Company)"], x["Role Pitched"], x["Location"], x["Email"],
            normalize_status(x["Status"], x["Date Contacted"], as_of), x["Status"],
            x["Original Message Summary"][:160], x["Notes"][:160], "Recruiter Tracker",
        ])
    for d, who, firm, client, role, verdict, status in data.get("old_2026_messages", []):
        stage = "Closed - I declined" if verdict == "Polite Decline" else "Closed - Went cold"
        log.append([as_date(d), who, firm, client, role, "", "", stage,
                    f"{verdict} / {status}", "", "", "2026 Messages sheet"])
    log.auto_filter.ref = f"A1:L{max(log.max_row, 2)}"
    for row in log.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
    add_list_validation(log, "H", rng["stage"], rows=2000)

    # ---- Recruiters: one row per person (the CRM) ----
    people = OrderedDict()
    for x in log_rows:
        name = re.sub(r"\s+", " ", x["Recruiter Name"]).strip()
        if not name:
            continue
        k = name.lower()
        p = people.setdefault(k, {"name": name, "agency": "", "email": "", "phone": "",
                                  "li": "", "roles": [], "clients": [], "dates": [],
                                  "last_status": "", "notes": ""})
        for fld, col in (("agency", "Company/Agency"), ("email", "Email"),
                         ("phone", "Phone"), ("li", "LinkedIn Profile URL")):
            v = x[col].strip()
            if v and not v.startswith("(") and "no-reply" not in v and "inmail-hit" not in v:
                p[fld] = p[fld] or v
        if x["Role Pitched"]:
            p["roles"].append(x["Role Pitched"])
        c = x["Recruiting For (Company)"].strip()
        if c and not c.startswith("(") and "not" not in c.lower():
            p["clients"].append(c)
        p["dates"].append(x["Date Contacted"])
        if x["Date Contacted"] >= max(p["dates"]):
            p["last_status"] = x["Status"]
    for name, agency, email, tier, note in data.get("email_only_recruiters", []):
        p = people.setdefault(name.lower(), {"name": name, "agency": agency, "email": email,
                                             "phone": "", "li": "", "roles": [], "clients": [],
                                             "dates": [], "last_status": "", "notes": ""})
        p["email"] = p["email"] or email
        p["agency"] = p["agency"] or agency
        p["tier"], p["notes"] = tier, note
    for pp in data.get("pipeline", []):
        for nm in re.split(r"[;/]", pp["recruiter"]):
            k = nm.strip().lower()
            if k in people:
                people[k]["dates"].append(pp["last"])

    tier_a = {n.lower() for n in data.get("tier_a_names", [])}
    dnc = {n.lower() for n in data.get("dnc_names", [])}
    headers = ["Recruiter", "Agency", "Tier", "Roles they pitch", "Clients seen", "Email",
               "Phone", "LinkedIn", "# Pitches", "First contact", "Last contact",
               "Next nurture", "Last status", "Notes"]
    widths = [24, 26, 6, 50, 30, 32, 14, 30, 9, 12, 12, 12, 30, 45]
    style_header(recs, headers, widths)
    rows = []
    for k, p in people.items():
        if "tier" in p:
            tier = p["tier"]
        elif k in dnc or "spam" in p["last_status"].lower():
            tier = "DNC"
        elif k in tier_a:
            tier = "A"
        elif any(TARGET_ROLE.search(r) for r in p["roles"]):
            tier = "B"
        else:
            tier = "C"
        ds = sorted(d for d in p["dates"] if d)
        rows.append((tier, ds[-1] if ds else "", p, ds))
    order = {t: i for i, t in enumerate(TIERS)}
    rows.sort(key=lambda t: t[1], reverse=True)  # newest first within each tier
    rows.sort(key=lambda t: order[t[0]])
    for i, (tier, lastd, p, ds) in enumerate(rows, start=2):
        roles = list(OrderedDict.fromkeys(p["roles"]))
        recs.append([
            p["name"], p["agency"], tier, "; ".join(roles[:4]),
            "; ".join(OrderedDict.fromkeys(p["clients"])), p["email"], p["phone"], p["li"],
            len(p["roles"]) or None, as_date(ds[0]) if ds else None,
            as_date(lastd) if lastd else None,
            f'=IF(K{i}="","",IF(C{i}="A",K{i}+42,IF(C{i}="B",K{i}+90,"")))',
            p["last_status"], p["notes"],
        ])
        for idx in (9, 10, 11):
            recs.cell(row=i, column=idx + 1).number_format = "yyyy-mm-dd"
    recs.conditional_formatting.add(
        "A2:N1000", FormulaRule(formula=['AND($L2<>"",$L2<=TODAY())'], fill=HOT_FILL))
    recs.conditional_formatting.add(
        "A2:N1000", FormulaRule(formula=['$C2="DNC"'], font=Font(color="C00000", strike=True)))
    add_list_validation(recs, "C", rng["tier"])
    recs.auto_filter.ref = "A1:N1000"

    # ---- Job Leads: postings found by sweeps ----
    headers = ["Date found", "Score", "Title", "Company", "Req ID", "Salary low",
               "Salary high", "Salary source", "Texas / remote eligibility", "Posted",
               "Status", "Days since posted", "Notes"]
    widths = [11, 7, 40, 24, 11, 11, 11, 16, 28, 11, 18, 10, 60]
    style_header(jobs, headers, widths)
    for i, j in enumerate(data.get("jobs", []), start=2):
        found, score, title, co, req, lo, hi, src, elig, posted, status, notes = j
        if status == "new" and (as_of - date.fromisoformat(posted)).days > 30:
            status = "Stale - recheck"
        jobs.append([as_date(found), score, title, co, req, lo or None, hi or None, src,
                     elig, as_date(posted), status, f'=IF(J{i}="","",TODAY()-J{i})', notes])
        for idx in (0, 9):
            jobs.cell(row=i, column=idx + 1).number_format = "yyyy-mm-dd"
        for idx in (5, 6):
            jobs.cell(row=i, column=idx + 1).number_format = "$#,##0"
    jobs.auto_filter.ref = "A1:M1000"

    # ---- Templates ----
    style_header(tmpl, ["Template", "Use when", "Text (edit the [brackets])"], [24, 34, 110])
    tmpl.freeze_panes = "A2"
    for t in TEMPLATES:
        tmpl.append(list(t))
    for row in tmpl.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP

    # ---- Start Here ----
    if not include_log:
        # The live Drive copy keeps its message log in the auto-updated tracker instead.
        wb.remove(log)
    build_start(start, data.get("pipeline", []), as_of, include_log, log_url)
    wb.save(out)
    return {"pipeline": len(data.get("pipeline", [])), "log": log.max_row - 1,
            "recruiters": recs.max_row - 1, "jobs": jobs.max_row - 1, "log_tab": include_log}


def build_start(ws, pipeline, as_of, include_log=True, log_url=""):
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 90
    ws["A1"] = "Damian - Job Search HQ"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Built {as_of.isoformat()}. Counts below update live; the action list is a snapshot."
    ws["A2"].font = Font(italic=True, color="808080")

    ws["A4"] = "Pipeline right now"
    ws["A4"].font = H2_FONT
    counts = [
        ("Reply owed by ME", '=COUNTIF(Pipeline!B:B,"2 Reply owed (me)")'),
        ("Waiting on recruiter", '=COUNTIF(Pipeline!B:B,"3 Waiting on recruiter")'),
        ("Submitted / RTR", '=COUNTIF(Pipeline!B:B,"4 Submitted / RTR")'),
        ("Interviewing", '=COUNTIF(Pipeline!B:B,"5 Interviewing")'),
        ("Offers", '=COUNTIF(Pipeline!B:B,"6 Offer")'),
        ("OVERDUE next actions",
         '=COUNTIFS(Pipeline!S:S,"<"&TODAY(),Pipeline!B:B,"<>Closed*")'),
        ("Recruiters due a nurture touch",
         '=COUNTIFS(Recruiters!L:L,"<="&TODAY(),Recruiters!L:L,"<>")'),
        ("A-tier recruiters", '=COUNTIF(Recruiters!C:C,"A")'),
    ]
    if include_log:
        counts.append(("Messages logged (all time)", "=COUNTA('Inbox Log'!A:A)-1"))
    for i, (label, f) in enumerate(counts, start=5):
        ws.cell(row=i, column=1, value=label)
        c = ws.cell(row=i, column=2, value=f)
        c.font = Font(bold=True, size=12)
    r = 5 + len(counts) + 1

    ws.cell(row=r, column=1, value=f"Do these first ({as_of.isoformat()})").font = H2_FONT
    r += 1
    rank = {"6 Offer": 0, "5 Interviewing": 1, "4 Submitted / RTR": 2, "2 Reply owed (me)": 3,
            "3 Waiting on recruiter": 4}
    todo = sorted((p for p in pipeline if p["stage"] in rank and p["priority"] != "Cold"),
                  key=lambda p: (rank[p["stage"]], PRIORITIES.index(p["priority"])))
    for n, p in enumerate(todo, start=1):
        ws.cell(row=r, column=1, value=f"{n}. {p['company']}")
        ws.cell(row=r, column=2, value=p["stage"])
        c = ws.cell(row=r, column=3, value=p["next"])
        c.alignment = WRAP
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="How this workbook works").font = H2_FONT
    r += 1
    for tab, what in [
        ("Pipeline", "One row per live opportunity. Only this tab drives your day. Keep Stage, Next action and Due filled in on every open row."),
        ("Recruiters", "One row per recruiter. Tier A = placed or submitted you, or a specialist in your field; B = pitches your target roles; C = off-target; DNC = never reply. Next nurture = A every 6 weeks, B every 90 days."),
        ("Inbox Log", ("" if include_log else
                       "Lives in the separate sheet 'LinkedIn Recruiter Tracker - Damian Ruiz', "
                       f"which your daily inbox run updates automatically: {log_url} ")
         + "Every recruiter message, including declines. You add to it but never delete "
         "from it. It answers 'has this person contacted me before?'."),
        ("Job Leads", "Postings you found yourself (HiringCafe/Jobright sweeps). When you act on one, add it to Pipeline."),
        ("Templates", "Reply templates, so a reply takes 2 minutes instead of sitting as a draft for 3 weeks."),
    ]:
        ws.cell(row=r, column=1, value=tab).font = Font(bold=True)
        c = ws.cell(row=r, column=3, value=what)
        c.alignment = WRAP
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="Rules").font = H2_FONT
    r += 1
    for rule in RULES:
        c = ws.cell(row=r, column=3, value=rule)
        c.alignment = WRAP
        ws.cell(row=r, column=1, value="-")
        r += 1


RULES = [
    "Reply to any recruiter within 1 business day, even if it's a decline. A draft that sits more than 24h goes to the top of the list.",
    "Remote only. Floor is $85K (= $41/hr W2). No Enbridge. These are the three auto-declines, so use the Templates tab.",
    "One agency per requisition. Before you sign an RTR, search Pipeline for the same role title and client. Being submitted twice gets you disqualified.",
    "After you send a resume or sign an RTR: nudge once after 5 business days, a second time at 10, then mark Closed - Went cold.",
    "After an interview: thank-you the same day, and nudge the recruiter at day 7.",
    "Every open Pipeline row needs a Due date. Red rows are overdue and are what you work first.",
    "A-tier recruiters get a short 'still looking: remote contracts/procurement, $85K+' note every 6 weeks, even when nothing is live.",
    "Friday, 20 minutes: clear red rows, move anything idle 14+ days forward or to Closed, and send nurture notes where Next nurture is due.",
]

TEMPLATES = [
    ("Remote-only decline", "Onsite or hybrid role.",
     "Hi [Name], thanks for thinking of me. I'm only considering fully remote roles right now, so I'll pass on this one. I focus on contracts, subcontracts and procurement at $85K+ (or $41+/hr W2), so please keep me in mind for remote roles in that space. Best, Damian"),
    ("Clarify before engaging", "Role looks right but work model, pay or client is missing.",
     "Hi [Name], thanks for reaching out. The [Role] scope fits my background in contracts and subcontract administration. Before we go further, could you confirm: (1) is it 100% remote, (2) the pay rate/range and W2 vs C2C, (3) the end client, and (4) contract length? Happy to set up a call once I have those. Best, Damian"),
    ("Yes - send resume / RTR", "Remote, $85K+, and not already submitted by another agency.",
     "Hi [Name], thanks. This fits well. Attached is my updated resume. I authorize [Agency] to represent me for [Role] at [Client] at [$rate]/hr W2. Please confirm I haven't already been submitted to this client for this req. Best, Damian"),
    ("Status nudge (5 business days)", "Resume or RTR sent and no update.",
     "Hi [Name], checking in on [Role] at [Client]. Any update from the hiring team? I'm still very interested. Thanks, Damian"),
    ("Post-interview nudge (day 7)", "Interviewed and no decision.",
     "Hi [Name], it's been about a week since my interview with [Client] for [Role]. Have they shared any feedback or timing? I enjoyed the conversation with [Interviewer] and remain very interested. Thanks, Damian"),
    ("Rejected - keep the door open", "Recruiter tells you it's a no.",
     "Hi [Name], thanks for letting me know, and for the opportunity. I'd love to be considered for the next [contracts / LNG / energy] search you run. Still focused on remote contracts and procurement roles. Best, Damian"),
    ("A-tier nurture (every 6 weeks)", "Next nurture date has arrived.",
     "Hi [Name], a quick update: I'm still open to remote Contracts / Subcontracts / Procurement roles, ideally $85K+ or $41+/hr W2. If anything comes across your desk, I'd appreciate a heads up. Hope things are going well. Damian"),
]


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--recruiters")
    ap.add_argument("--data")
    ap.add_argument("--out", required=True)
    ap.add_argument("--as-of", default=date.today().isoformat())
    ap.add_argument("--no-log", action="store_true",
                    help="omit the Inbox Log tab (Drive copy links to the live tracker instead)")
    ap.add_argument("--log-url", default="")
    a = ap.parse_args()
    print(build(a.recruiters, a.data, a.out, date.fromisoformat(a.as_of),
                include_log=not a.no_log, log_url=a.log_url))
